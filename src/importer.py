"""📥 ایمپورت محتوا از CSV / XLSX / JSON / Google Sheet."""
from __future__ import annotations

import csv
import io
import json
import re


class ImportError_(Exception):
    """خطای قابل‌خواندن برای ادمین."""
    pass


def _sheet_to_csv_url(url: str) -> str | None:
    """تبدیل لینک گوگل شیت به لینک خروجی CSV."""
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url)
    if not m:
        return None
    gid_m = re.search(r"[#&?]gid=(\d+)", url)
    gid = gid_m.group(1) if gid_m else "0"
    return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv&gid={gid}"


async def fetch_bytes(url: str) -> bytes:
    import aiohttp  # lazy import — فقط موقع نیاز
    csv_url = _sheet_to_csv_url(url)
    if csv_url:
        url = csv_url
    async with aiohttp.ClientSession() as sess:
        async with sess.get(url, timeout=aiohttp.ClientTimeout(total=30)) as resp:
            if resp.status != 200:
                raise ImportError_(f"😕 دانلود نشد (HTTP {resp.status}) — برای گوگل‌ شیت مطمئن شو «Anyone with the link» فعاله.")
            return await resp.read()


def parse_questions_csv(raw: bytes) -> list[dict]:
    """CSV با ستون‌های: question, option1, option2, option3, option4, correct, explanation"""
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for i, r in enumerate(reader, start=2):
        try:
            q = {
                "text": (r.get("question") or r.get("text") or "").strip(),
                "opt1": (r.get("option1") or r.get("opt1") or "").strip(),
                "opt2": (r.get("option2") or r.get("opt2") or "").strip(),
                "opt3": (r.get("option3") or r.get("opt3") or "").strip(),
                "opt4": (r.get("option4") or r.get("opt4") or "").strip(),
                "correct": int((r.get("correct") or "").strip()),
                "explanation": (r.get("explanation") or "").strip(),
            }
        except (ValueError, AttributeError):
            raise ImportError_(f"⚠️ ردیف {i} مشکل داره — ستون‌ها رو چک کن.")
        if not all([q["text"], q["opt1"], q["opt2"], q["opt3"], q["opt4"]]):
            raise ImportError_(f"⚠️ ردیف {i} ناقصه — همهٔ ستون‌ها باید پر باشن.")
        if not (1 <= q["correct"] <= 4):
            raise ImportError_(f"⚠️ ردیف {i}: مقدار correct باید بین ۱ تا ۴ باشه.")
        rows.append(q)
    if not rows:
        raise ImportError_("⚠️ هیچ ردیفی توی فایل پیدا نشد!")
    return rows


def parse_cards_csv(raw: bytes) -> list[dict]:
    """CSV با ستون‌های: front, back, example, tip"""
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for i, r in enumerate(reader, start=2):
        c = {
            "front": (r.get("front") or "").strip(),
            "back": (r.get("back") or "").strip(),
            "example": (r.get("example") or "").strip(),
            "tip": (r.get("tip") or "").strip(),
        }
        if not c["front"] or not c["back"]:
            raise ImportError_(f"⚠️ ردیف {i}: front و back اجباری‌ان.")
        rows.append(c)
    if not rows:
        raise ImportError_("⚠️ هیچ ردیفی توی فایل پیدا نشد!")
    return rows


def parse_questions_json(raw: bytes) -> list[dict]:
    try:
        data = json.loads(raw.decode("utf-8"))
    except Exception:
        raise ImportError_("⚠️ JSON معتبر نیست!")
    if not isinstance(data, list):
        raise ImportError_("⚠️ JSON باید یه لیست از سؤال‌ها باشه.")
    rows = []
    for i, r in enumerate(data, start=1):
        opts = r.get("options") or [r.get("opt1"), r.get("opt2"), r.get("opt3"), r.get("opt4")]
        if len(opts) != 4 or not all(opts):
            raise ImportError_(f"⚠️ سؤال {i} — باید دقیقاً ۴ گزینه داشته باشه.")
        correct = r.get("correct")
        if isinstance(correct, str):
            correct = int(correct)
        rows.append({
            "text": str(r.get("text") or r.get("question") or ""),
            "opt1": str(opts[0]), "opt2": str(opts[1]),
            "opt3": str(opts[2]), "opt4": str(opts[3]),
            "correct": correct, "explanation": str(r.get("explanation") or ""),
        })
    return parse_questions_csv(
        ("question,option1,option2,option3,option4,correct,explanation\n" +
         "\n".join(",".join([
             _csv_escape(r["text"]), _csv_escape(r["opt1"]), _csv_escape(r["opt2"]),
             _csv_escape(r["opt3"]), _csv_escape(r["opt4"]), str(r["correct"]),
             _csv_escape(r["explanation"]),
         ]) for r in rows)).encode()
    )


def parse_cards_json(raw: bytes) -> list[dict]:
    try:
        data = json.loads(raw.decode("utf-8"))
    except Exception:
        raise ImportError_("⚠️ JSON معتبر نیست!")
    rows = []
    for r in data:
        rows.append({
            "front": str(r.get("front") or ""), "back": str(r.get("back") or ""),
            "example": str(r.get("example") or ""), "tip": str(r.get("tip") or ""),
        })
    return parse_cards_csv(
        ("front,back,example,tip\n" + "\n".join(
            ",".join(_csv_escape(x) for x in [r["front"], r["back"], r["example"], r["tip"]])
            for r in rows)).encode()
    )


def _csv_escape(s: str) -> str:
    s = str(s)
    return '"' + s.replace('"', '""') + '"' if ("," in s or '"' in s or "\n" in s) else s


def parse_questions_xlsx(raw: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError_("⚠️ برای اکسل باید openpyxl نصب بشه — از requirements.txt نصب کن.")
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ImportError_("⚠️ اکسل حداقل باید یه ردیف سربرگ + یه ردیف داده داشته باشه.")
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    body = []
    for r in rows[1:]:
        body.append(dict(zip(header, ["" if v is None else str(v) for v in r])))
    csv_text = "question,option1,option2,option3,option4,correct,explanation\n"
    for r in body:
        csv_text += ",".join([
            _csv_escape(r.get("question", "") or r.get("text", "")),
            _csv_escape(r.get("option1", "") or r.get("opt1", "")),
            _csv_escape(r.get("option2", "") or r.get("opt2", "")),
            _csv_escape(r.get("option3", "") or r.get("opt3", "")),
            _csv_escape(r.get("option4", "") or r.get("opt4", "")),
            str(r.get("correct", "")),
            _csv_escape(r.get("explanation", "")),
        ]) + "\n"
    return parse_questions_csv(csv_text.encode())


def parse_cards_xlsx(raw: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError_("⚠️ برای اکسل باید openpyxl نصب بشه.")
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ImportError_("⚠️ اکسل حداقل باید یه ردیف سربرگ + یه ردیف داده داشته باشه.")
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    body = []
    for r in rows[1:]:
        body.append(dict(zip(header, ["" if v is None else str(v) for v in r])))
    csv_text = "front,back,example,tip\n"
    for r in body:
        csv_text += ",".join([
            _csv_escape(r.get("front", "")), _csv_escape(r.get("back", "")),
            _csv_escape(r.get("example", "")), _csv_escape(r.get("tip", "")),
        ]) + "\n"
    return parse_cards_csv(csv_text.encode())
